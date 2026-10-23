"""Birleşik Markdown taslağından görselleri içeren Word ve PDF üretir.

Gereksinimler: Python 3, python-docx, pandoc, XeLaTeX, DejaVu ve Liberation yazı tipleri.
Görsel dosyalarının içeriği değiştirilmez; yalnız belge içindeki boyutu ayarlanır.
"""

from datetime import datetime, timezone
from pathlib import Path
import os
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
STEM = "ARASTIRMA_ONERISI_BIRLESIK_TASLAK"


def make_reference(path, sunum=False, butce=False):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    section.header_distance = section.footer_distance = Cm(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Arial" if sunum else "Times New Roman"
    normal.font.size = Pt(10 if sunum else 11)
    normal.paragraph_format.line_spacing = 1.05 if sunum else 1.12
    normal.paragraph_format.space_after = Pt(4 if sunum else 6)
    normal.paragraph_format.widow_control = True
    for name, size in [("Heading 1", 22 if sunum else 17), ("Heading 2", 12 if sunum else 14), ("Heading 3", 10.5 if sunum else 11.5)]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("183A57")
        style.paragraph_format.space_before = Pt(9 if sunum else 12)
        style.paragraph_format.space_after = Pt(4 if sunum else 6)
        style.paragraph_format.keep_with_next = True
    caption = document.styles["Caption"]
    caption.font.name = "Arial" if sunum else "Times New Roman"
    caption.font.size = Pt(8.5 if sunum else 9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string("475569")
    header = section.header.paragraphs[0]
    header.text = "TÜBİTAK 2209-A  |  Araştırma önerisi" + ("" if sunum else " taslağı")
    if butce:
        header.text = "TÜBİTAK 2209-A | Gerekçeli bütçe"
    if sunum:
        header.text = "2209-A ÜNİVERSİTE ÖĞRENCİLERİ ARAŞTIRMA PROJELERİ DESTEKLEME PROGRAMI\nARAŞTIRMA ÖNERİSİ"
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.runs[0].font.bold = True
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("64748B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT if sunum else WD_ALIGN_PARAGRAPH.CENTER
    if not sunum:
        footer.add_run("Ahmet Talha Arslan  ·  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(path)


def finish_layout(path, title=None, sunum=False):
    document = Document(path)
    properties = document.core_properties
    properties.title = title or "İKA — TÜBİTAK 2209-A Araştırma Önerisi, Birleşik Taslak"
    properties.author = "Ahmet Talha Arslan"
    properties.subject = "Araştırma önerisi; yöntem, ön çalışma ve deney planı" if sunum else "Planlanan araştırma; deney sonuçları içermez"
    properties.modified = datetime.now(timezone.utc)

    def paragraph_rule(paragraph, edge, color):
        borders = OxmlElement("w:pBdr")
        rule = OxmlElement("w:" + edge)
        for key, value in [("val", "single"), ("sz", "5"), ("space", "10"), ("color", color)]:
            rule.set(qn("w:" + key), value)
        borders.append(rule)
        paragraph._p.get_or_add_pPr().append(borders)

    if sunum:
        document.sections[0].different_first_page_header_footer = True
        paragraph_rule(document.sections[0].header.paragraphs[0], "bottom", "879CAE")
        paragraph_rule(document.paragraphs[0], "top", "203E5C")
        document.paragraphs[0].paragraph_format.space_before = Pt(20)
        cover = True
        for paragraph in document.paragraphs:
            if paragraph.text == "A. GENEL BİLGİLER":
                cover = False
                paragraph.paragraph_format.page_break_before = True
            if cover and paragraph.text.strip():
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.space_after = Pt(14)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.color.rgb = RGBColor.from_string("203E5C")
                if paragraph.text == "TÜBİTAK 2209-A":
                    paragraph.paragraph_format.space_before = Pt(15)
                    for run in paragraph.runs:
                        run.font.size = Pt(22)
                elif paragraph.text.startswith("Kullanıcı "):
                    paragraph.paragraph_format.space_before = Pt(30)
                    paragraph.paragraph_format.space_after = Pt(28)
                    for run in paragraph.runs:
                        run.font.size = Pt(14)
                elif paragraph.text.startswith("ÜNİVERSİTE ÖĞRENCİLERİ"):
                    for run in paragraph.runs:
                        run.font.size = Pt(15)
                elif paragraph.text == "ARAŞTIRMA ÖNERİSİ":
                    paragraph.paragraph_format.space_before = Pt(18)
                    for run in paragraph.runs:
                        run.font.size = Pt(19)
                elif paragraph.text == "2026 YILI":
                    paragraph.paragraph_format.space_before = Pt(30)
                    paragraph.paragraph_format.keep_with_next = False
                    for run in paragraph.runs:
                        run.font.size = Pt(14)
                    paragraph_rule(paragraph, "bottom", "CBD7E0")

    for paragraph in document.paragraphs:
        if sunum and paragraph.text in {"KAYNAKLAR", "6. EKLER"}:
            paragraph.paragraph_format.page_break_before = True
        if not sunum and paragraph.text in {
            "1. Bilimsel nitelik", "2. Yöntem", "3. Proje yönetimi",
            "4. Araştırma önerisinin yaygın etkisi", "Kaynakça",
            "Son düzenlemede tamamlanacak alanlar — hazırlık notu",
            "1. ARAŞTIRMA ÖNERİSİNİN BİLİMSEL NİTELİĞİ", "2. YÖNTEM",
            "3. PROJE YÖNETİMİ", "4. ARAŞTIRMA ÖNERİSİNİN YAYGIN ETKİSİ",
            "KAYNAKLAR",
        }:
            paragraph.paragraph_format.page_break_before = True
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(2)
        if paragraph.text.startswith(("Şekil 1.", "Şekil 2.", "Şekil 3.", "Şekil 4.")):
            paragraph.style = document.styles["Caption"]
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False

    for shape in document.inline_shapes:
        factor = min(Cm(16.8) / shape.width, Cm(17.6) / shape.height, 1)
        shape.width = int(shape.width * factor)
        shape.height = int(shape.height * factor)

    if sunum:
        logo = document.inline_shapes[0]
        ratio = logo.height / logo.width
        logo.width = Cm(2.6)
        logo.height = int(logo.width * ratio)
        # Özetin tamamını tek, düzenlenebilir kutuda tut.
        paragraphs = list(document.paragraphs)
        start = next(i for i, p in enumerate(paragraphs) if p.text == "ÖZET")
        end = next(i for i, p in enumerate(paragraphs) if p.text.startswith("1. ARAŞTIRMA"))
        summary = document.add_table(rows=1, cols=1)
        paragraphs[start]._p.addnext(summary._tbl)
        cell = summary.cell(0, 0)
        for paragraph in paragraphs[start + 1:end]:
            cell._tc.append(paragraph._p)
        empty = cell.paragraphs[0]
        if not empty.text:
            cell._tc.remove(empty._p)

    for table in document.tables:
        table.style = "Table Grid"
        table.autofit = False
        count = len(table.columns)
        label = table.cell(0, 0).text
        if count == 2:
            widths = [4.2, 12.8]
        elif count == 3 and label == "Koşul":
            widths = [1.5, 12.8, 2.7]
        elif count == 3 and label == "Kişisel kaynak":
            widths = [3.3, 6.0, 7.7]
        elif count == 3 and label == "Bileşen":
            widths = [3.7, 1.4, 11.9]
        elif count == 3 and label == "Risk":
            widths = [4.1, 11.4, 1.5]
        elif count == 3 and label == "Ölçüt":
            widths = [3.7, 7.3, 6.0]
        elif count == 3 and label == "Çıktı kategorisi":
            widths = [3.4, 6.3, 7.3]
        elif count == 3 and label == "Senaryo":
            widths = [1.7, 7.8, 7.5]
        elif count == 3:
            widths = [3.0, 5.6, 8.4]
        elif count == 4 and label == "İP":
            widths = [0.8, 10.0, 4.6, 1.6]
        elif count == 4 and label == "Kalem":
            widths = [8.5, 1.7, 3.4, 3.4]
        elif count == 4 and label == "Koşul":
            widths = [1.1, 3.3, 11.1, 1.5]
        elif count == 5 and label == "Dönem":
            widths = [1.5, 3.6, 2.4, 8.1, 1.4]
        elif count == 4:
            widths = [3.2, 4.1, 8.1, 1.6]
        else:
            widths = [17 / count] * count
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        summary_table = sunum and count == 1
        borders = OxmlElement("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            rule = OxmlElement("w:" + edge)
            rule.set(qn("w:val"), "single")
            rule.set(qn("w:sz"), "4")
            rule.set(qn("w:color"), "879CAE")
            borders.append(rule)
        table._tbl.tblPr.append(borders)
        if not summary_table:
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row_index, row in enumerate(table.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                if row_index == 0 and not summary_table:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EAF1F7")
                    cell._tc.get_or_add_tcPr().append(shade)
                for paragraph in cell.paragraphs:
                    if label == "Dönem":
                        paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.size = Pt(10 if summary_table else 9.5)
                        if row_index == 0 and not summary_table:
                            run.font.bold = True
    document.save(path)


def main(stem=STEM, title=None, sunum=False):
    source = ROOT / f"{stem}.md"
    output = ROOT / f"{stem}.docx"
    with tempfile.TemporaryDirectory(prefix="ika-belge-") as temporary:
        reference = Path(temporary) / "reference.docx"
        butce = stem == "BASVURU_BUTCE_ONERISI"
        make_reference(reference, sunum, butce)
        subprocess.run([
            "pandoc", str(source), "--from=markdown-implicit_figures+autolink_bare_uris", "--standalone",
            "--metadata=lang:tr-TR", f"--resource-path={ROOT}",
            f"--reference-doc={reference}", "--output", str(output),
        ], check=True, cwd=ROOT)
        finish_layout(output, title, sunum)
        pdf_directory = Path(temporary) / "pdf"
        pdf_directory.mkdir()
        cache_directory = Path(temporary) / "cache"
        cache_directory.mkdir()
        generated_pdf = pdf_directory / f"{stem}.pdf"
        subprocess.run([
            "pandoc", str(source), "--from=markdown-implicit_figures+autolink_bare_uris", "--standalone",
            "--pdf-engine=xelatex", f"--resource-path={ROOT}",
            f"--lua-filter={ROOT / 'belge_kaynaklari/pdf_yerlesim.lua'}",
            f"--include-in-header={ROOT / ('belge_kaynaklari/pdf_sunum_baslik.tex' if sunum else 'belge_kaynaklari/pdf_butce_baslik.tex' if butce else 'belge_kaynaklari/pdf_baslik.tex')}",
            "-V", "mainfont=" + ("Liberation Sans" if sunum else "DejaVu Serif"),
            "-V", "sansfont=" + ("Liberation Sans" if sunum else "DejaVu Sans"),
            "-V", "papersize=a4",
            "-V", "geometry=margin=20mm", "-V", "fontsize=" + ("10pt" if sunum else "11pt"),
            "-V", "linestretch=" + ("1.05" if sunum else "1.08"), "-V", "colorlinks=true",
            "--output", str(generated_pdf),
        ], check=True, cwd=ROOT, env={**os.environ, "XDG_CACHE_HOME": str(cache_directory)})
        if not generated_pdf.is_file() or not generated_pdf.stat().st_size:
            raise RuntimeError("PDF çıktısı oluşmadı.")
        shutil.copy2(generated_pdf, ROOT / generated_pdf.name)
    pdf = ROOT / f"{stem}.pdf"
    output.chmod(0o644)
    pdf.chmod(0o644)
    print(f"Word: {output.name}\nPDF: {pdf.name}")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stem", default=STEM, help="Kaynak Markdown dosyasının uzantısız adı")
    parser.add_argument("--title", help="Word belgesinin özelliklerindeki başlık")
    parser.add_argument("--sunum", action="store_true", help="Kapaklı danışman sunum düzeni")
    arguments = parser.parse_args()
    main(arguments.stem, arguments.title, arguments.sunum)
